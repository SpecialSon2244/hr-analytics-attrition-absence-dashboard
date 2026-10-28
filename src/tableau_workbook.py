"""Build tableau/hr_analytics_dashboard.xlsx (and the same tables as CSV in exports/)
from data/hr.duckdb.

Every sheet is a small, tidy table with rates already computed as percentages
(0-100), so charts in Tableau Public web authoring need no calculated fields.
Sheets whose name ends in _SIMULATED come from the SIMULATED recruiting pipeline.
"""
import pathlib

import duckdb
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / "tableau" / "hr_analytics_dashboard.xlsx"
EXPORTS = ROOT / "exports"

RATE = "round(100.0 * sum(f.is_leaver) / count(*), 1)"

DRIVERS = [  # (factor label, dim_employee column, order expression)
    ("Overtime", "overtime", "CASE overtime WHEN 'No' THEN 1 ELSE 2 END"),
    ("Years at company", "tenure_band", "tenure_band"),
    ("Monthly income", "income_band", "income_band"),
    ("Business travel", "business_travel",
     "CASE business_travel WHEN 'No travel' THEN 1 WHEN 'Travels rarely' THEN 2 ELSE 3 END"),
    ("Age band", "age_band", "age_band"),
    ("Job level", "'Level ' || job_level", "job_level"),
    ("Job satisfaction", "job_satisfaction",
     "CASE job_satisfaction WHEN 'Low' THEN 1 WHEN 'Medium' THEN 2 WHEN 'High' THEN 3 ELSE 4 END"),
]

drivers_sql = "\nUNION ALL\n".join(
    f"""SELECT '{label}' AS "Factor", {i} AS "Factor order",
               regexp_replace(({col})::VARCHAR, '^[0-9]\\. ', '') AS "Group",
               dense_rank() OVER (ORDER BY {order}) AS "Group order",
               count(*) AS "Headcount", sum(f.is_leaver) AS "Leavers", {RATE} AS "Attrition rate %"
        FROM fact_attrition_snapshot f JOIN dim_employee e USING (employee_key)
        GROUP BY ({col}), ({order})"""
    for i, (label, col, order) in enumerate(DRIVERS, 1))

SHEETS = {
    "kpis": None,  # built in kpis_sql()
    "attrition_by_role": f"""
        SELECT e.department_name AS "Department", e.job_role AS "Job role",
               count(*) AS "Headcount", sum(f.is_active) AS "Active", sum(f.is_leaver) AS "Leavers",
               {RATE} AS "Attrition rate %",
               round(avg(f.tenure_years), 1) AS "Avg tenure (years)",
               (SELECT round(100.0 * sum(is_leaver) / count(*), 1) FROM fact_attrition_snapshot) AS "Company attrition rate %"
        FROM fact_attrition_snapshot f JOIN dim_employee e USING (employee_key)
        GROUP BY ALL ORDER BY "Attrition rate %" DESC""",
    "attrition_drivers": f"""SELECT * FROM ({drivers_sql}) ORDER BY "Factor order", "Group order" """,
    "overtime_income": f"""
        SELECT 'Overtime: ' || e.overtime AS "Overtime",
               regexp_replace(e.income_band, '^[0-9]\\. ', '') AS "Monthly income band",
               left(e.income_band, 1)::INT AS "Band order",
               count(*) AS "Headcount", sum(f.is_leaver) AS "Leavers", {RATE} AS "Attrition rate %"
        FROM fact_attrition_snapshot f JOIN dim_employee e USING (employee_key)
        GROUP BY ALL ORDER BY 1, 3""",
    "absence_reason": """
        SELECT r.reason_name AS "Reason", r.reason_group AS "Reason group", count(*) AS "Records",
               sum(a.absence_hours) AS "Absence hours",
               round(100.0 * sum(a.absence_hours) / (SELECT sum(absence_hours) FROM fact_absence), 1) AS "Share of hours %",
               round(avg(a.absence_hours), 1) AS "Avg hours per record"
        FROM fact_absence a JOIN dim_absence_reason r USING (reason_code)
        GROUP BY ALL HAVING sum(a.absence_hours) > 0 ORDER BY "Absence hours" DESC""",
    "absence_calendar": """
        SELECT m.month_number AS "Month number", m.month_name AS "Month",
               w.iso_weekday AS "Weekday number", w.weekday_name AS "Weekday",
               count(a.absence_id) AS "Records", coalesce(sum(a.absence_hours), 0) AS "Absence hours"
        FROM (SELECT DISTINCT month_number, month_name FROM fact_absence WHERE month_number BETWEEN 1 AND 12) m
        CROSS JOIN (SELECT DISTINCT iso_weekday, weekday_name FROM fact_absence) w
        LEFT JOIN fact_absence a ON a.month_number = m.month_number AND a.iso_weekday = w.iso_weekday
        GROUP BY ALL ORDER BY 1, 3""",
    "hiring_funnel_SIMULATED": """
        WITH s AS (
            SELECT 'All departments' AS dept, stage, stage_order, count(*) AS n FROM fact_candidate_stage GROUP BY ALL
            UNION ALL
            SELECT d.department_name, stage, stage_order, count(*) FROM fact_candidate_stage
            JOIN dim_department d USING (department_key) GROUP BY ALL)
        SELECT dept AS "Department", regexp_replace(stage, '^[0-9]\\. ', '') AS "Stage", stage_order AS "Stage order",
               n AS "Candidates",
               round(100.0 * n / max(n) OVER (PARTITION BY dept), 1) AS "% of applicants",
               round(100.0 * n / lag(n) OVER (PARTITION BY dept ORDER BY stage_order), 1) AS "Conversion from previous stage %",
               'SIMULATED' AS "Data"
        FROM s ORDER BY CASE dept WHEN 'All departments' THEN 0 ELSE 1 END, 1, 3""",
    "time_to_hire_SIMULATED": """
        SELECT r.requisition_id AS "Requisition", d.department_name AS "Department", r.job_role AS "Job role",
               r.opened_date AS "Opened", r.closed_date AS "Filled",
               year(r.closed_date) || '-Q' || quarter(r.closed_date) AS "Quarter filled",
               r.time_to_hire_days AS "Time to hire (days)", r.time_to_fill_days AS "Time to fill (days)",
               r.hired_source AS "Hire source", r.applicants AS "Applicants", 'SIMULATED' AS "Data"
        FROM fact_requisition r JOIN dim_department d USING (department_key)
        WHERE r.status = 'Filled' ORDER BY r.closed_date, 1""",
    "hiring_dept_SIMULATED": """
        WITH o AS (
            SELECT department_key, sum((offer_outcome = 'Accepted')::INT) acc,
                   sum((offer_outcome IN ('Accepted', 'Declined'))::INT) dec
            FROM fact_candidate GROUP BY 1)
        SELECT d.department_name AS "Department", count(*) AS "Requisitions",
               sum((status = 'Filled')::INT) AS "Filled", sum((status = 'Open')::INT) AS "Open",
               sum((status = 'Cancelled')::INT) AS "Cancelled",
               median(time_to_hire_days) AS "Median time to hire (days)",
               median(time_to_fill_days) AS "Median time to fill (days)",
               round(100.0 * any_value(o.acc) / any_value(o.dec), 1) AS "Offer acceptance %",
               'SIMULATED' AS "Data"
        FROM fact_requisition r JOIN dim_department d USING (department_key) JOIN o USING (department_key)
        GROUP BY 1 ORDER BY 2 DESC""",
}


def kpis_sql() -> str:
    return """
        WITH w AS (SELECT count(*) hc, sum(is_active) act, sum(is_leaver) lv, avg(tenure_years) ten
                   FROM fact_attrition_snapshot),
             a AS (SELECT sum(absence_hours) h, count(DISTINCT absence_employee_key) emp FROM fact_absence),
             r AS (SELECT median(time_to_hire_days) tth, median(time_to_fill_days) ttf FROM fact_requisition),
             o AS (SELECT 100.0 * sum((offer_outcome = 'Accepted')::INT)
                          / sum((offer_outcome IN ('Accepted', 'Declined'))::INT) oa FROM fact_candidate)
        SELECT * FROM (VALUES
          (1, 'Headcount', 'IBM (fictional)', (SELECT hc FROM w)::DOUBLE, (SELECT format('{:,}', hc) FROM w)),
          (2, 'Attrition rate', 'IBM (fictional)', (SELECT round(100.0 * lv / hc, 1) FROM w), (SELECT round(100.0 * lv / hc, 1) || '%' FROM w)),
          (3, 'Avg tenure', 'IBM (fictional)', (SELECT round(ten, 1) FROM w), (SELECT round(ten, 1) || ' yrs' FROM w)),
          (4, 'Absence hours', 'UCI (real, other company)', (SELECT h FROM a)::DOUBLE, (SELECT format('{:,}', h) FROM a)),
          (5, 'Absence hours per employee per year', 'UCI (real, other company)',
              (SELECT round(h / emp / 3.0, 1) FROM a), (SELECT round(h / emp / 3.0, 1)::VARCHAR FROM a)),
          (6, 'Median time to hire', 'SIMULATED', (SELECT tth FROM r)::DOUBLE, (SELECT round(tth)::INT || ' days' FROM r)),
          (7, 'Median time to fill', 'SIMULATED', (SELECT ttf FROM r)::DOUBLE, (SELECT round(ttf)::INT || ' days' FROM r)),
          (8, 'Offer acceptance', 'SIMULATED', (SELECT round(oa, 1) FROM o), (SELECT round(oa, 1) || '%' FROM o))
        ) t("KPI order", "KPI", "Source", "Value", "Display")"""


ABOUT = [
    ("About this file", ""),
    ("Project", "HR Analytics: Attrition, Absence & Time to Hire"),
    ("IBM sheets", "kpis (rows 1-3), attrition_by_role, attrition_drivers, overtime_income. IBM HR Analytics Employee "
                   "Attrition & Performance: a FICTIONAL dataset created by IBM data scientists, 1,470 employees, "
                   "one undated snapshot. ODbL / DbCL."),
    ("UCI sheets", "kpis (rows 4-5), absence_reason, absence_calendar. UCI Absenteeism at work (id 445): real records "
                   "from a courier company in Brazil, Jul 2007 to Jul 2010, 36 employees, 740 records. CC BY 4.0. "
                   "A different company from IBM: never join it to the IBM sheets."),
    ("SIMULATED sheets", "kpis (rows 6-8) and every sheet ending in _SIMULATED. Generated by src/simulate_hiring.py "
                         "(seed 20261028), using IBM departments and job roles. Not real recruiting data."),
    ("Attrition rate %", "Leavers (Attrition = Yes) / headcount in the snapshot x 100"),
    ("Time to hire", "Days from the hired candidate's application to the offer acceptance (SIMULATED)"),
    ("Time to fill", "Days from requisition opened to offer acceptance (SIMULATED)"),
    ("Offer acceptance %", "Accepted offers / (accepted + declined) offers (SIMULATED)"),
    ("Percentages", "Stored as 0-100 numbers, not fractions"),
    ("Built by", "src/tableau_workbook.py from data/hr.duckdb"),
]


def main() -> None:
    con = duckdb.connect(str(ROOT / "data" / "hr.duckdb"), read_only=True)
    EXPORTS.mkdir(exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name, sql in SHEETS.items():
        rel = con.sql(sql if sql else kpis_sql())
        rel.write_csv(str(EXPORTS / f"{name}.csv"))
        cols, rows = rel.columns, rel.fetchall()
        ws = wb.create_sheet(name)
        ws.append(cols)
        for r in rows:
            ws.append([float(v) if hasattr(v, "is_finite") else v for v in r])  # Decimal -> float
        for i, c in enumerate(cols, 1):
            ws.cell(row=1, column=i).font = Font(bold=True)
            ws.column_dimensions[get_column_letter(i)].width = max(12, len(c) + 2)
            if c in ("Opened", "Filled"):
                for cell in ws[get_column_letter(i)][1:]:
                    cell.number_format = "yyyy-mm-dd"
        ws.freeze_panes = "A2"
        print(f"{name}: {len(rows)} rows")
    ws = wb.create_sheet("about")
    for r in ABOUT:
        ws.append(list(r))
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 120
    wb.save(OUT)
    print(f"{OUT} {OUT.stat().st_size / 1e6:.2f} MB")


if __name__ == "__main__":
    main()
