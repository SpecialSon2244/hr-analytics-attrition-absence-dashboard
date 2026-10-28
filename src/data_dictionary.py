"""Write docs/data_dictionary.md: every model table and column with its type (read from
data/hr.duckdb), source and meaning, plus the columns of the Tableau / CSV exports."""
import pathlib

import duckdb
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]

TABLES = {
    "dim_department": ("Shared (IBM departments)", "One row per IBM department. Used by IBM and SIMULATED tables."),
    "dim_job_role": ("Shared (IBM job roles)", "One row per department + job role (Manager exists in all three departments)."),
    "dim_date": ("Calendar", "One row per day, 2025-01-01 to 2026-12-31. Used by the SIMULATED recruiting facts only."),
    "dim_employee": ("IBM (fictional)", "One row per IBM employee, with attributes and derived bands."),
    "fact_attrition_snapshot": ("IBM (fictional)", "One row per IBM employee at the undated snapshot. Grain: employee."),
    "dim_absence_employee": ("UCI (real)", "One row per courier-company employee (36). Not the same people as IBM."),
    "dim_absence_reason": ("UCI (real)", "Absence reason codes: 1-21 ICD-10 chapters, 22-28 other reasons, 0 none recorded."),
    "fact_absence": ("UCI (real)", "One row per absence record (740), Jul 2007 to Jul 2010. No year column in the source."),
    "fact_requisition": ("SIMULATED", "One row per SIMULATED requisition (300). Grain: requisition."),
    "fact_candidate": ("SIMULATED", "One row per SIMULATED candidate application. Grain: candidate x requisition."),
    "fact_candidate_stage": ("SIMULATED", "One row per SIMULATED candidate per stage reached, with the stage date."),
}

COLS = {
    "department_key": "Surrogate key of the department", "department_name": "Department (IBM: Sales, Research & Development, Human Resources)",
    "job_role_key": "Surrogate key of department + job role", "job_role": "Job role as in the IBM file",
    "date_key": "Date as yyyymmdd integer", "date": "Calendar date", "year": "Year", "quarter": "Quarter number",
    "quarter_name": "Q1 to Q4", "month": "Month number", "month_name": "Month abbreviation", "year_month": "yyyy-mm",
    "iso_weekday": "ISO weekday, 1 = Monday", "weekday_name": "Weekday abbreviation", "is_weekend": "Saturday or Sunday",
    "employee_key": "IBM EmployeeNumber", "job_level": "Job level 1 (lowest) to 5", "age": "Age in years",
    "age_band": "Age band", "gender": "Gender as in the file", "marital_status": "Marital status",
    "education_level": "Education 1-5 mapped to Below college, College, Bachelor, Master, Doctor",
    "education_field": "Field of education", "business_travel": "No travel / Travels rarely / Travels frequently",
    "distance_from_home": "Distance from home (units not documented)", "overtime": "Works overtime (Yes/No)",
    "monthly_income": "Monthly income (no currency given in the dataset)", "income_band": "Monthly income band (sortable prefix)",
    "stock_option_level": "Stock option level 0-3", "percent_salary_hike": "Last salary increase, %",
    "performance_rating": "Performance rating (only 3 and 4 occur)", "job_satisfaction": "Job satisfaction 1-4 mapped to Low..Very high",
    "environment_satisfaction": "Environment satisfaction 1-4 mapped to Low..Very high",
    "work_life_balance": "Work-life balance 1-4 mapped to Bad, Good, Better, Best", "num_companies_worked": "Number of previous employers",
    "total_working_years": "Total working years", "training_times_last_year": "Trainings last year",
    "years_at_company": "Years at the company (tenure)", "tenure_band": "Tenure band (sortable prefix)",
    "years_in_current_role": "Years in current role", "years_since_last_promotion": "Years since last promotion",
    "years_with_current_manager": "Years with current manager",
    "headcount_at_start": "Always 1: the employee is counted in the snapshot headcount",
    "is_leaver": "1 if Attrition = Yes (left). The dataset does not say when or whether it was voluntary",
    "is_active": "1 if Attrition = No", "termination_type": "Always 'Not recorded' (the source has no reason for leaving)",
    "tenure_years": "Years at company", "absence_employee_key": "UCI employee ID as 'UCI-nn'",
    "min_age_in_records": "Age (constant per employee in the file)", "service_time_years": "Service time in years",
    "distance_to_work_km": "Distance from residence to work, km", "transportation_expense": "Transportation expense (currency not stated)",
    "children": "Number of children ('Son' in the source)", "social_drinker": "Social drinker flag", "social_smoker": "Social smoker flag",
    "pets": "Number of pets", "body_mass_index": "Body mass index", "reason_code": "UCI reason code 0-28",
    "reason_name": "Reason label (short ICD-10 chapter name or the UCI label)", "reason_group": "ICD disease / Non-ICD / No reason recorded",
    "absence_id": "Row number of the record in the source file", "month_number": "Month of absence 1-12 (0 = not recorded, 3 rows)",
    "season_uci_code_label": "UCI season code as labelled by UCI (does not match the months; not used in charts)",
    "season_by_month": "Season derived from the month, Southern Hemisphere", "absence_hours": "Absenteeism time in hours",
    "absence_days_8h": "absence_hours / 8 (assumes an 8-hour day; used for the Bradford factor)",
    "is_absence_spell": "1 if the record has more than 0 hours", "disciplinary_failure": "Disciplinary failure flag (these rows have 0 hours)",
    "work_load_average_per_day": "Work load average per day (as in source)", "hit_target": "Hit target % (as in source)",
    "age_at_absence": "Age", "candidate_id": "SIMULATED candidate ID", "requisition_id": "SIMULATED requisition ID",
    "source": "SIMULATED applicant source", "applied_date_key": "Application date as yyyymmdd", "applied_date": "Application date",
    "screened_date": "Date the candidate passed screening (blank if not)", "interview_date": "First interview date",
    "final_interview_date": "Final interview date", "offer_date": "Offer date", "offer_decision_date": "Date the offer was accepted or declined",
    "offer_outcome": "Accepted / Declined / Pending / No offer", "start_date": "Planned start date after notice period",
    "furthest_stage": "Furthest stage reached", "time_to_hire_days": "Days from application to offer acceptance (hired candidates only)",
    "is_simulated": "Always true: SIMULATED row", "stage": "Stage with sortable prefix", "stage_order": "Stage order 1-6",
    "stage_date_key": "Stage date as yyyymmdd", "stage_date": "Date the stage was reached", "status": "Filled / Open / Cancelled at 30 Sep 2026",
    "opened_date_key": "Requisition opened date as yyyymmdd", "opened_date": "Requisition opened date",
    "closed_date_key": "Filled or cancelled date as yyyymmdd", "closed_date": "Filled (offer accepted) or cancelled date",
    "hired_candidate_id": "Candidate who accepted", "hired_source": "Source of the hired candidate",
    "time_to_fill_days": "Days from requisition opened to offer acceptance", "applicants": "Applicants on the requisition",
    "offers": "Offers made on the requisition",
}


def main() -> None:
    con = duckdb.connect(str(ROOT / "data" / "hr.duckdb"), read_only=True)
    md = ["# Data dictionary", "", "Generated by `src/data_dictionary.py` from `data/hr.duckdb`.", "",
          "**Three separate subject areas.** IBM = fictional IBM HR dataset. UCI = real absence records of a "
          "different company (Brazilian courier). SIMULATED = generated recruiting pipeline. The UCI tables never "
          "join to IBM or SIMULATED tables.", "",
          "Power BI copies of these tables are in `powerbi/model/` (SIMULATED tables have `_SIMULATED` in the file "
          "name; `dim_employee` there drops the department columns, which come from `dim_job_role`; two small "
          "helper dimensions `dim_absence_month` and `dim_weekday` are added).", ""]
    missing = []
    for t, (src, desc) in TABLES.items():
        md += [f"## {t}", "", f"Source: **{src}**. {desc} Rows: {con.execute(f'SELECT count(*) FROM {t}').fetchone()[0]:,}.", "",
               "| Column | Type | Description |", "|---|---|---|"]
        for c, ty in con.execute(f"SELECT column_name, data_type FROM information_schema.columns WHERE table_name = '{t}' ORDER BY ordinal_position").fetchall():
            if c not in COLS:
                missing.append(f"{t}.{c}")
            md.append(f"| {c} | {ty} | {COLS.get(c, '')} |")
        md.append("")
    md += ["## Tableau workbook and CSV exports", "",
           "`tableau/hr_analytics_dashboard.xlsx` has one sheet per chart; `exports/<sheet>.csv` holds the same tables. "
           "Percentages are 0-100. Sheets ending in `_SIMULATED` are SIMULATED.", "",
           "| Sheet | Rows | Columns |", "|---|---|---|"]
    wb = load_workbook(ROOT / "tableau" / "hr_analytics_dashboard.xlsx", read_only=True)
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(values_only=True))
        md.append(f"| {name} | {len(rows) - 1} | {', '.join(str(c) for c in rows[0])} |")
    (ROOT / "docs" / "data_dictionary.md").write_text("\n".join(md) + "\n")
    if missing:
        raise SystemExit(f"columns without description: {missing}")
    print("docs/data_dictionary.md written")


if __name__ == "__main__":
    main()
